"""Fill a .docx template from document state, and render the same template as
an editable HTML preview.

Tokens (placed in the template by tools/build_*_template.py or by hand in Word):

    {{field}}          inline text field
    {{sig:role}}       signatory name: config.json signatories, overridable per document
    {{annex:field}}    "see Annex B" / "N/A" marker for an attachment field
    {{#field}}         paragraph replaced by the field's block (paragraphs, table or images)
    {{#annexes}}       paragraph replaced by one annex page per provided attachment

A token must sit inside a single run to keep its formatting; if Word splits it
across runs the paragraph is merged into its first run so it still fills.
"""
from __future__ import annotations

import copy
import html
import io
import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.image.image import Image as DocxImage
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from docx.table import Table
from docx.text.paragraph import Paragraph

from .registry import Template, is_placeholder, resolved_values

TOKEN = re.compile(r"\{\{(#?)(\w+)(?::(\w+))?\}\}")
BLANK_LINE = "_" * 45
FONT = "Aptos"


class Context:
    def __init__(self, tpl: Template, state: dict, config: dict, upload_dir: Path):
        self.tpl = tpl
        self.values = resolved_values(tpl, state, config)
        self.skipped = set((state or {}).get("skipped") or [])
        sigs = dict(config.get("signatories", {}))
        for role, name in ((state or {}).get("signatories") or {}).items():
            if name and str(name).strip():
                sigs[role] = str(name).strip()
        self.signatories = sigs
        self.upload_dir = upload_dir
        self.annexes = []  # (letter, field)
        for f in tpl.fields:
            if f.get("attachment") and self.values.get(f["id"]):
                self.annexes.append((chr(ord("A") + len(self.annexes)), f))
        self.annex_letter = {f["id"]: letter for letter, f in self.annexes}

    def annex_marker(self, fid: str) -> tuple[str, str]:
        if fid in self.annex_letter:
            return f"– see Annex {self.annex_letter[fid]}", "attached"
        if fid in self.skipped:
            return "– N/A", "na"
        return "", "pending"

    def image_path(self, upload_id: str) -> Path | None:
        if not re.fullmatch(r"[\w-]+", upload_id or ""):
            return None
        hits = list(self.upload_dir.glob(upload_id + ".*"))
        return hits[0] if hits else None


def _paragraphs_of(container):
    for p in container.paragraphs:
        yield p
    for t in getattr(container, "tables", []):
        for row in t.rows:
            for cell in row.cells:
                yield from _paragraphs_of(cell)


def _all_paragraphs(doc):
    yield from _paragraphs_of(doc)
    for s in doc.sections:
        for part in (s.header, s.footer, s.first_page_header, s.first_page_footer):
            if part is not None and not part.is_linked_to_previous:
                yield from _paragraphs_of(part)


# ================================================================ DOCX


def fill_docx(tpl: Template, state: dict, config: dict, upload_dir: Path) -> bytes:
    ctx = Context(tpl, state, config, upload_dir)
    doc = Document(str(tpl.path))
    blocks = []
    for p in list(_all_paragraphs(doc)):
        m = TOKEN.fullmatch(p.text.strip())
        if m and m.group(1) == "#":
            blocks.append((p, m.group(2)))
        elif "{{" in p.text:
            _fill_inline(p, ctx)
    for p, name in blocks:
        if name == "annexes":
            _docx_annexes(doc, p, ctx)
        else:
            f = tpl.field(name)
            if f:
                _docx_field_block(doc, p, f, ctx.values.get(name), ctx, in_annex=False)
        p._p.getparent().remove(p._p)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _inline_text(m: re.Match, ctx: Context) -> str:
    _, name, sub = m.groups()
    if name == "sig":
        return ctx.signatories.get(sub, f"[{sub.upper()}]")
    if name == "annex":
        return ctx.annex_marker(sub)[0]
    v = ctx.values.get(name)
    return v if isinstance(v, str) and v else BLANK_LINE


def _fill_inline(p: Paragraph, ctx: Context):
    for r in p.runs:
        if "{{" in r.text:
            r.text = TOKEN.sub(lambda m: _inline_text(m, ctx), r.text)
    if TOKEN.search(p.text) and p.runs:  # token split across runs
        text = TOKEN.sub(lambda m: _inline_text(m, ctx), p.text)
        p.runs[0].text = text
        for r in p.runs[1:]:
            r._r.getparent().remove(r._r)


def _style_run(run, size=12, bold=None, italic=None, color=None):
    run.font.name = FONT
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def _new_par(doc, anchor: Paragraph, text="", *, model: Paragraph | None = None, size=12,
             bold=None, italic=None, align=None, after=None, page_break=False) -> Paragraph:
    """Insert a paragraph before anchor. With model, copy its paragraph and run formatting."""
    if model is not None:
        el = copy.deepcopy(model._p)
        for child in list(el):
            if child.tag != qn("w:pPr"):
                el.remove(child)
        anchor._p.addprevious(el)
        p = Paragraph(el, anchor._parent)
        rpr = model.runs[0]._r.rPr if model.runs else None
        r = p.add_run(text)
        if rpr is not None:
            r._r.insert(0, copy.deepcopy(rpr))
    else:
        p = doc.add_paragraph()
        anchor._p.addprevious(p._p)
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(text)
        _style_run(r, size, bold, italic)
    if align is not None:
        p.alignment = align
    if after is not None:
        p.paragraph_format.space_after = Pt(after)
    if page_break:
        p.paragraph_format.page_break_before = True
    return p


def _docx_field_block(doc, anchor: Paragraph, f: dict, value, ctx: Context, in_annex: bool):
    t = f["type"]
    if t in ("longtext", "text"):
        paras = [s.strip() for s in (value or "").split("\n") if s.strip()]
        if not paras:
            paras = [BLANK_LINE] * 3
        for text in paras:
            if in_annex:
                p = _new_par(doc, anchor, text, align=WD_ALIGN_PARAGRAPH.JUSTIFY, after=8)
                p.paragraph_format.first_line_indent = Inches(0.5)
            else:
                _new_par(doc, anchor, text, model=anchor, after=6 if len(paras) > 1 else None)
    elif t == "table":
        _docx_table(doc, anchor, f, value or [])
    elif t == "images":
        _docx_images(doc, anchor, f, value or [], ctx)


def _shade(cell, hex_fill):
    tcpr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    tcpr.append(shd)


def _docx_table(doc, anchor: Paragraph, f: dict, rows: list):
    cols = list(f.get("columns", []))
    numbered = f.get("numbered")
    if numbered:
        cols = ["No."] + cols
    table = doc.add_table(rows=1, cols=len(cols))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, c in enumerate(cols):
        cell = table.rows[0].cells[i]
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        _style_run(cell.paragraphs[0].add_run(c), 11, bold=True)
        _shade(cell, "D9DEEA")
    for n, row in enumerate(rows, 1):
        cells = table.add_row().cells
        vals = ([str(n)] if numbered else []) + list(row)
        for i, v in enumerate(vals):
            _style_run(cells[i].paragraphs[0].add_run(v), 11)
            if numbered and i == 0:
                cells[i].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    table.autofit = False
    total, first = 6.2, (0.55 if numbered else None)
    rest = (total - (first or 0)) / max(1, len(cols) - (1 if numbered else 0))
    widths = [Inches(first if numbered and i == 0 else rest) for i in range(len(cols))]
    for i, col in enumerate(table.columns):
        col.width = widths[i]
    for row in table.rows:
        for i, cell in enumerate(row.cells):
            cell.width = widths[i]
    tr = table.rows[0]._tr  # repeat header row on each page
    trpr = tr.get_or_add_trPr()
    hdr = OxmlElement("w:tblHeader")
    trpr.append(hdr)
    anchor._p.addprevious(table._tbl)


def _docx_images(doc, anchor: Paragraph, f: dict, images: list, ctx: Context):
    max_w, max_h = 6.0, (3.8 if f["id"] == "pictures" else 8.0)
    for n, item in enumerate(images, 1):
        path = ctx.image_path(item["id"])
        p = _new_par(doc, anchor, "", align=WD_ALIGN_PARAGRAPH.CENTER, after=4)
        p.paragraph_format.keep_with_next = True
        if path is None:
            _style_run(p.add_run("[image not found]"), 11, italic=True)
        else:
            try:
                img = DocxImage.from_file(str(path))
                w_in = img.px_width / (img.horz_dpi or 96)
                h_in = img.px_height / (img.vert_dpi or 96)
                scale = min(max_w / w_in, max_h / h_in)
                p.add_run().add_picture(str(path), width=Inches(w_in * scale))
            except Exception:
                _style_run(p.add_run("[unsupported image format]"), 11, italic=True)
        caption = item.get("caption") or ""
        if f["id"] == "pictures":
            caption = f"Figure {n}. {caption}".strip()
        if caption:
            _new_par(doc, anchor, caption, size=11, italic=True, align=WD_ALIGN_PARAGRAPH.CENTER, after=14)
        else:
            _new_par(doc, anchor, "", size=6, after=10)


def _docx_annexes(doc, anchor: Paragraph, ctx: Context):
    for letter, f in ctx.annexes:
        _new_par(doc, anchor, f"ANNEX {letter}", size=14, bold=True,
                 align=WD_ALIGN_PARAGRAPH.CENTER, page_break=True)
        _new_par(doc, anchor, f.get("annex_title", f["label"]).upper(), size=12, bold=True,
                 align=WD_ALIGN_PARAGRAPH.CENTER, after=4)
        event = ctx.values.get("event_name")
        if event:
            _new_par(doc, anchor, event, size=11, italic=True, align=WD_ALIGN_PARAGRAPH.CENTER, after=16)
        else:
            _new_par(doc, anchor, "", size=6, after=12)
        _docx_field_block(doc, anchor, f, ctx.values.get(f["id"]), ctx, in_annex=True)


# ================================================================ HTML preview

E = html.escape
TWIP_PX = 1 / 15  # 1440 twips per inch, 96 px per inch


def _twips(v) -> float:
    return int(v) * TWIP_PX if v not in (None, "") else 0.0


class _Numbering:
    """Minimal list numbering so "1.", "4.2" etc. show in the preview."""

    def __init__(self, doc):
        self.levels = {}  # (numId, ilvl) -> (fmt, text, left, hanging, isLgl)
        self.counters = {}
        try:
            root = doc.part.numbering_part.element
        except Exception:
            return
        abstract = {a.get(qn("w:abstractNumId")): a for a in root.findall(qn("w:abstractNum"))}
        for num in root.findall(qn("w:num")):
            a = abstract.get(num.find(qn("w:abstractNumId")).get(qn("w:val")))
            if a is None:
                continue
            for lvl in a.findall(qn("w:lvl")):
                g = lambda tag: (lvl.find(qn(tag)).get(qn("w:val")) if lvl.find(qn(tag)) is not None else None)
                ind = lvl.find(qn("w:pPr") + "/" + qn("w:ind"))
                self.levels[(num.get(qn("w:numId")), int(lvl.get(qn("w:ilvl"))))] = (
                    g("w:numFmt") or "decimal", g("w:lvlText") or "",
                    _twips(ind.get(qn("w:left")) if ind is not None else 0),
                    _twips(ind.get(qn("w:hanging")) if ind is not None else 0),
                    lvl.find(qn("w:isLgl")) is not None)

    @staticmethod
    def _fmt(n, fmt):
        if fmt == "lowerLetter":
            return chr(96 + n)
        if fmt == "upperLetter":
            return chr(64 + n)
        if fmt in ("lowerRoman", "upperRoman"):
            vals = [(10, "x"), (9, "ix"), (5, "v"), (4, "iv"), (1, "i")]
            out = ""
            for v, s in vals:
                while n >= v:
                    out, n = out + s, n - v
            return out if fmt == "lowerRoman" else out.upper()
        if fmt == "bullet":
            return "•"
        return str(n)

    def label(self, p: Paragraph):
        ppr = p._p.pPr
        numpr = ppr.find(qn("w:numPr")) if ppr is not None else None
        if numpr is None or numpr.find(qn("w:numId")) is None:
            return None
        num_id = numpr.find(qn("w:numId")).get(qn("w:val"))
        ilvl_el = numpr.find(qn("w:ilvl"))
        ilvl = int(ilvl_el.get(qn("w:val"))) if ilvl_el is not None else 0
        info = self.levels.get((num_id, ilvl))
        if not info:
            return None
        fmt, text, left, hanging, is_lgl = info
        c = self.counters.setdefault(num_id, [0] * 9)
        c[ilvl] += 1
        for k in range(ilvl + 1, 9):
            c[k] = 0
        def sub(m):
            k = int(m.group(1)) - 1
            lf = self.levels.get((num_id, k), ("decimal",))[0]
            return self._fmt(max(c[k], 1), "decimal" if is_lgl and lf != "bullet" else lf)
        return re.sub(r"%(\d)", sub, text) if fmt != "bullet" else "•", left, hanging


def _par_style(p: Paragraph, num_info=None) -> tuple[str, list]:
    css = []
    ppr = p._p.pPr
    jc = ppr.find(qn("w:jc")).get(qn("w:val")) if ppr is not None and ppr.find(qn("w:jc")) is not None else None
    css.append({"center": "text-align:center", "right": "text-align:right", "end": "text-align:right",
                "both": "text-align:justify"}.get(jc, "text-align:left"))
    left = hanging = first = 0.0
    if num_info:
        left, hanging = num_info[1], num_info[2]
    elif p.style is not None and p.style.paragraph_format.left_indent:
        left = p.style.paragraph_format.left_indent / 9525  # EMU -> px
    ind = ppr.find(qn("w:ind")) if ppr is not None else None
    if ind is not None:
        if ind.get(qn("w:left")) is not None or ind.get(qn("w:start")) is not None:
            left = _twips(ind.get(qn("w:left")) or ind.get(qn("w:start")))
        if ind.get(qn("w:hanging")) is not None:
            hanging = _twips(ind.get(qn("w:hanging")))
        if ind.get(qn("w:firstLine")) is not None:
            first = _twips(ind.get(qn("w:firstLine")))
    css.append(f"padding-left:{left:.1f}px")
    if first:
        css.append(f"text-indent:{first:.1f}px")
    after = p.paragraph_format.space_after
    if after is None:
        after = p.style.paragraph_format.space_after if p.style is not None else None
    after_px = (after / 9525) if after is not None else 160 * TWIP_PX
    css.append(f"margin:0 0 {after_px:.1f}px 0")
    if p.paragraph_format.page_break_before:
        css.append("break-before:page")
    tabs = []
    tabs_el = ppr.find(qn("w:tabs")) if ppr is not None else None
    if tabs_el is not None:
        tabs = [round(_twips(t.get(qn("w:pos")))) for t in tabs_el.findall(qn("w:tab"))
                if t.get(qn("w:val")) not in ("clear",)]
    return ";".join(css), tabs


def _run_css(r) -> str:
    css = []
    if r.bold:
        css.append("font-weight:700")
    if r.italic:
        css.append("font-style:italic")
    if r.underline:
        css.append("text-decoration:underline")
    if r.font.size:
        css.append(f"font-size:{r.font.size.pt * 4 / 3:.1f}px")
    return ";".join(css)


def _html_inline_token(m: re.Match, ctx: Context) -> str:
    _, name, sub = m.groups()
    if name == "sig":
        role_label = ctx.tpl.schema.get("signatories", {}).get(sub, sub)
        v = ctx.signatories.get(sub, "")
        cls = "fld sig" + (" placeholder" if is_placeholder(v) or not v else "")
        return (f'<span class="{cls}" contenteditable="plaintext-only" spellcheck="false" '
                f'data-sig="{E(sub)}" data-label="{E(role_label)}" title="{E(role_label)} (click to edit)">{E(v)}</span>')
    if name == "annex":
        text, state = ctx.annex_marker(sub)
        if state == "pending":
            text = "– not yet provided"
        return f'<span class="annex-tag {state}" data-annex="{E(sub)}">{E(text)}</span>'
    f = ctx.tpl.field(name)
    if not f:
        return E(m.group(0))
    v = ctx.values.get(name) or ""
    return (f'<span class="fld{"" if v else " empty"}" contenteditable="plaintext-only" '
            f'data-field="{E(name)}" data-label="{E(f["label"])}">{E(v)}</span>')


def _html_runs(p: Paragraph, ctx: Context) -> str:
    out = []
    for r in p.runs:
        if r._r.rPr is not None and r._r.rPr.find(qn("w:vanish")) is not None:
            continue
        text = r.text
        if not text:
            if r._r.find(qn("w:br")) is not None:
                out.append("<br>")
            continue
        css = _run_css(r)
        parts = []
        for seg in re.split(r"(\t|\n)", text):
            if seg == "\t":
                parts.append('<span class="tab"></span>')
            elif seg == "\n":
                parts.append("<br>")
            else:
                pos, buf = 0, []
                for m in TOKEN.finditer(seg):
                    buf.append(E(seg[pos:m.start()]))
                    buf.append(_html_inline_token(m, ctx))
                    pos = m.end()
                buf.append(E(seg[pos:]))
                parts.append("".join(buf))
        inner = "".join(parts)
        out.append(f'<span style="{css}">{inner}</span>' if css else inner)
    return "".join(out)


def _html_par(p: Paragraph, ctx: Context, numbering: _Numbering) -> str:
    num = numbering.label(p)
    css, tabs = _par_style(p, num)
    label = ""
    if num:
        size = _run_css(p.runs[0]) if p.runs else ""
        label = (f'<span class="num" style="min-width:{num[2]:.1f}px;margin-left:-{num[2]:.1f}px;{size}">'
                 f'{E(num[0])}</span>')
    body = _html_runs(p, ctx)
    if not body.strip() and not label:
        body = "&nbsp;"
    data = f' data-tabs="{",".join(map(str, tabs))}"' if tabs else ""
    return f'<p class="dp" style="{css}"{data}>{label}{body}</p>'


def _html_field_block(f: dict, value, ctx: Context, model_css: str = "") -> str:
    fid = E(f["id"])
    t = f["type"]
    if t in ("longtext", "text"):
        paras = [s.strip() for s in (value or "").split("\n") if s.strip()]
        inner = "".join(f"<p>{E(s)}</p>" for s in paras)
        return (f'<div class="fld blk{"" if paras else " empty"}" contenteditable="true" data-field="{fid}" '
                f'data-kind="longtext" data-label="{E(f["label"])}" style="{model_css}">{inner}</div>')
    if t == "table":
        cols = list(f.get("columns", []))
        num = bool(f.get("numbered"))
        head = ("<th class='no'>No.</th>" if num else "") + "".join(f"<th>{E(c)}</th>" for c in cols)
        rows = []
        for i, row in enumerate(value or []):
            cells = (f"<td class='no'>{i + 1}</td>" if num else "") + "".join(
                f'<td contenteditable="plaintext-only" data-field="{fid}" data-row="{i}" data-col="{j}">{E(c)}</td>'
                for j, c in enumerate(row))
            rows.append(f'<tr>{cells}<td class="row-x no-print"><button data-del-row="{i}" data-field="{fid}" '
                        f'title="Remove row">×</button></td></tr>')
        return (f'<table class="dtable" data-table="{fid}"><thead><tr>{head}<th class="row-x no-print"></th></tr></thead>'
                f'<tbody>{"".join(rows)}</tbody></table>'
                f'<button class="add-row no-print" data-add-row="{fid}">+ Add row</button>')
    if t == "images":
        figs = []
        for i, item in enumerate(value or []):
            prefix = f"Figure {i + 1}. " if f["id"] == "pictures" else ""
            figs.append(
                f'<figure class="dimg"><img src="/api/uploads/{E(item["id"])}" alt="">'
                f'<figcaption>{E(prefix)}<span class="fld cap{"" if item.get("caption") else " empty"}" '
                f'contenteditable="plaintext-only" data-field="{fid}" data-img="{i}" data-label="Add a caption">'
                f'{E(item.get("caption") or "")}</span></figcaption>'
                f'<button class="img-x no-print" data-del-img="{i}" data-field="{fid}" title="Remove image">×</button>'
                f'</figure>')
        return "".join(figs)
    return ""


def render_html(tpl: Template, state: dict, config: dict, upload_dir: Path) -> str:
    ctx = Context(tpl, state, config, upload_dir)
    doc = Document(str(tpl.path))
    numbering = _Numbering(doc)
    sec = doc.sections[0]
    page_css = (f"width:{sec.page_width / 9525:.0f}px;min-height:{sec.page_height / 9525:.0f}px;"
                f"padding:{sec.top_margin / 9525:.0f}px {sec.right_margin / 9525:.0f}px "
                f"{sec.bottom_margin / 9525:.0f}px {sec.left_margin / 9525:.0f}px")
    header = ""
    if sec.header is not None and not sec.header.is_linked_to_previous:
        header = "".join(_html_par(p, ctx, numbering) for p in sec.header.paragraphs)
        header = f'<div class="dhead" style="top:{(sec.header_distance or 0) / 9525:.0f}px">{header}</div>'

    body, annex_html = [], ""
    for el in doc.element.body.iterchildren():
        if el.tag == qn("w:p"):
            p = Paragraph(el, doc)
            m = TOKEN.fullmatch(p.text.strip())
            if m and m.group(1) == "#":
                name = m.group(2)
                if name == "annexes":
                    annex_html = _html_annexes(ctx, page_css)
                elif tpl.field(name):
                    css, _ = _par_style(p)
                    body.append(_html_field_block(tpl.field(name), ctx.values.get(name), ctx, css))
                continue
            body.append(_html_par(p, ctx, numbering))
        elif el.tag == qn("w:tbl"):
            t = Table(el, doc)
            rows = "".join("<tr>" + "".join(
                "<td>" + "".join(_html_par(cp, ctx, numbering) for cp in c.paragraphs) + "</td>"
                for c in r.cells) + "</tr>" for r in t.rows)
            body.append(f'<table class="dtable plain">{rows}</table>')

    return (f'<section class="page" style="{page_css}">{header}{"".join(body)}</section>{annex_html}')


def _html_annexes(ctx: Context, page_css: str) -> str:
    out = []
    event = ctx.values.get("event_name") or ""
    for letter, f in ctx.annexes:
        out.append(
            f'<section class="page annex" style="{page_css}">'
            f'<h2 class="annex-h">ANNEX {letter}</h2><h3 class="annex-t">{E(f.get("annex_title", f["label"]).upper())}</h3>'
            f'{f"<p class=annex-e>{E(event)}</p>" if event else ""}'
            f'{_html_field_block(f, ctx.values.get(f["id"]), ctx, "text-align:justify;text-indent:48px")}'
            f'</section>')
    return "".join(out)
