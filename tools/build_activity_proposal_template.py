"""Turn the blank ASCEnD Activity Proposal (source_forms/ACTIVITY-PROPOSAL.docx)
into the fillable template used by the generator.

The original layout, header/footer art, table styling and fonts are kept.
Instruction lines and blank spaces are replaced with tokens that app/engine.py
understands:

    {{field}}          inline text field
    {{#field}}         whole paragraph replaced by a block (paragraphs or images)
    {{rows:field}}     table row repeated once per row of a table field
    {{total:field}}    sum of a table field's Amount column
    {{sig:role}}       signatory name (config.json, editable per document)
    {{title:role}}     signatory position (schema label, editable per document)

Run from the project root:  python tools/build_activity_proposal_template.py
"""
import copy
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
SRC = ROOT / "source_forms" / "ACTIVITY-PROPOSAL.docx"
OUT = ROOT / "templates" / "activity_proposal" / "template.docx"


def expect(text, starts):
    if not text.strip().upper().startswith(starts.upper()):
        sys.exit(f"Unexpected template layout: wanted {starts!r}, found {text!r}")


def set_text(p, text):
    """Put text in the paragraph's first run (keeping its formatting) and drop the rest."""
    runs = p.runs
    if not runs:
        p.add_run(text)
        return p.runs[0]
    runs[0].text = text
    for r in runs[1:]:
        r._r.getparent().remove(r._r)
    return runs[0]


def cell_token(cell, text, bold=None, align=None):
    """Write text into an empty table cell using the cell's own run format (its paragraph mark)."""
    p = cell.paragraphs[0]
    r = p.add_run(text)
    mark = p._p.pPr.find(qn("w:rPr")) if p._p.pPr is not None else None
    if mark is not None:
        rpr = copy.deepcopy(mark)
        r._r.insert(0, rpr)
    if bold is not None:
        r.bold = bold
    if align is not None:
        p.alignment = align
    return r


def remove(el):
    el.getparent().remove(el)


def main():
    doc = Document(SRC)
    ps = list(doc.paragraphs)
    t_head, t_overview, t_budget, t_flow, t_approval = doc.tables

    for i, starts in ((3, "INTRODUCTION"), (4, "Brief Background"), (6, "GENERAL OBJECTIVES"),
                      (9, "NU CORE VALUES"), (10, "Indicate the NU Core"), (14, "TARGET SUSTAINABLE"),
                      (15, "Indicate the Sustainable"), (19, "ACTIVITY OVERVIEW"), (22, "BUDGET BREAKDOWN"),
                      (23, "Source of Budget"), (24, "Remaining Budget"), (27, "PROGRAMME FLOW"),
                      (30, "EVENT DETAILS"), (33, "APPROVAL")):
        expect(ps[i].text, starts)
    expect(t_head.cell(0, 0).text, "TITLE OF EVENT")
    expect(t_head.cell(1, 2).text, "DATE OF EVENT")
    expect(t_flow.cell(0, 0).text, "TITLE:")
    expect(t_approval.cell(1, 0).paragraphs[2].text, "Name of student")
    if len(t_overview.rows) != 7 or len(t_budget.rows) != 5 or len(t_flow.rows) != 3:
        sys.exit("Unexpected table sizes in the Activity Proposal")

    # a body paragraph in the form's own text style (Aptos 11, indented under the heading)
    proto = copy.deepcopy(ps[4]._p)
    from docx.text.paragraph import Paragraph

    def block_after(anchor_el, token, align=WD_ALIGN_PARAGRAPH.JUSTIFY):
        el = copy.deepcopy(proto)
        anchor_el.addnext(el)
        p = Paragraph(el, ps[4]._parent)
        r = set_text(p, token)
        r.italic = False
        r.bold = False
        p.alignment = align
        return el

    # header table: title, organization, date
    cell_token(t_head.cell(0, 1), "{{event_name}}")
    set_text(t_head.cell(1, 1).paragraphs[0], "{{organization_name}}")
    cell_token(t_head.cell(1, 3), "{{event_date}}")

    # I-IV: instruction lines become the written sections; one blank line is kept after each
    block_after(ps[3]._p, "{{#introduction}}")
    block_after(ps[6]._p, "{{#objectives}}")
    block_after(ps[9]._p, "{{#core_values}}")
    block_after(ps[14]._p, "{{#sdgs}}")
    for i in (4, 7, 10, 11, 12, 15, 16, 17):
        remove(ps[i]._p)

    # V: activity overview, filled in place
    cell_token(t_overview.cell(1, 0), "{{rows:activity_overview}}")
    for tr in t_overview._tbl.tr_lst[2:]:
        remove(tr)

    # VI: budget source / remaining (value not bold), items filled in place, TOTAL row added up
    for p, field in ((ps[23], "budget_source"), (ps[24], "remaining_budget")):
        colon = p.runs[1]
        value = copy.deepcopy(colon._r)
        colon._r.addnext(value)
        colon.text = ": "
        v = p.runs[2]
        v.text = "{{%s}}" % field
        v.bold = False
    cell_token(t_budget.cell(1, 0), "{{rows:budget_items}}")
    total_row = t_budget.rows[4]
    cell_token(total_row.cells[0], "TOTAL", bold=True, align=WD_ALIGN_PARAGRAPH.RIGHT)
    cell_token(total_row.cells[-1], "{{total:budget_items}}", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    for tr in t_budget._tbl.tr_lst[2:4]:
        remove(tr)

    # VII: programme flow title + rows
    title_p = t_flow.cell(0, 0).paragraphs[0]
    run = copy.deepcopy(title_p.runs[0]._r)
    title_p.runs[0]._r.addnext(run)
    title_p.runs[1].text = "{{program_title}}"
    cell_token(t_flow.cell(2, 0), "{{rows:program_flow}}")

    # VIII: event details text, then floor plan / layout images; the blank line before IX stays
    details = block_after(ps[30]._p, "{{#event_details}}")
    block_after(details, "{{#event_layout}}", align=WD_ALIGN_PARAGRAPH.CENTER)
    remove(ps[31]._p)

    # IX: printed names and positions over the signature space
    for col, role in ((0, "president"), (1, "adviser_1")):
        cps = t_approval.cell(1, col).paragraphs
        set_text(cps[2], "{{sig:%s}}" % role).bold = True
        set_text(cps[3], "{{title:%s}}" % role)
        set_text(cps[4], "{{organization_name}}")

    # keep each section heading with what follows it, and the approval block together
    for p in doc.paragraphs:
        pPr = p._p.pPr
        if pPr is not None and pPr.find(qn("w:numPr")) is not None:
            p.paragraph_format.keep_with_next = True
    ps[34].paragraph_format.keep_with_next = True

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print("wrote", OUT)


if __name__ == "__main__":
    main()
