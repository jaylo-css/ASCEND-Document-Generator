"""Build the fillable SDAO Financial Report template (ICB-SDA-2026-007).

SDAO gave this form only as a PDF (source_forms/FINANCIAL-REPORT.pdf), so it is
rebuilt here with python-docx: the layout, sizes and colors were measured from
the PDF, and the letterhead and "FOR INCUBATION ONLY" watermark are the PDF's own
images (extracted to tools/assets/financial_report/).

Tokens used (see app/engine.py):

    {{field}}          inline text field
    {{rows:field}}     table row repeated once per row of a table field
    {{total:field}}    sum of a table field's Amount column
    {{sig:role}}       signatory name (config.json, editable per document)
    {{title:role}}     signatory position (schema label, editable per document)

Run from the project root:  python tools/build_financial_report_template.py
"""
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Emu, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
ASSETS = ROOT / "tools" / "assets" / "financial_report"
OUT = ROOT / "templates" / "financial_report" / "template.docx"

FONT = "Aptos"
NAVY = "002060"
BOX_BLUE = "156082"
LEFT, RIGHT = 76.6, 67.4            # page margins in pt (table edges in the PDF)
CONTENT = 595.32 - LEFT - RIGHT     # 451.3 pt


def twips(pt):
    return str(round(pt * 20))


def style_run(run, size=9, bold=None, italic=None, underline=None, color=None):
    run.font.name = FONT
    run._r.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), FONT)
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if underline is not None:
        run.underline = underline
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    return run


def para(container, text="", size=9, bold=None, italic=None, underline=None, align=None, after=0, before=0,
         color=None, p=None):
    p = p or container.add_paragraph()
    pf = p.paragraph_format
    pf.space_after, pf.space_before, pf.line_spacing = Pt(after), Pt(before), 1.0
    if align is not None:
        p.alignment = align
    if text:
        style_run(p.add_run(text), size, bold, italic, underline, color)
    else:  # size the empty line itself
        mark = OxmlElement("w:rPr")
        sz = OxmlElement("w:sz")
        sz.set(qn("w:val"), str(int(size * 2)))
        mark.append(sz)
        p._p.get_or_add_pPr().append(mark)
    return p


def float_picture(paragraph, path, x_pt, y_pt, w_pt, h_pt, name, behind=False, z=1):
    """Picture anchored to the page at (x, y), like the PDF's letterhead and watermark."""
    inline = paragraph.add_run().add_picture(str(path), width=Pt(w_pt), height=Pt(h_pt))._inline
    anchor = parse_xml(
        f'<wp:anchor {nsdecls("wp", "a", "pic", "r")} distT="0" distB="0" distL="0" distR="0" simplePos="0" '
        f'relativeHeight="{z}" behindDoc="{1 if behind else 0}" locked="0" layoutInCell="1" allowOverlap="1">'
        f'<wp:simplePos x="0" y="0"/>'
        f'<wp:positionH relativeFrom="page"><wp:posOffset>{Emu(Pt(x_pt))}</wp:posOffset></wp:positionH>'
        f'<wp:positionV relativeFrom="page"><wp:posOffset>{Emu(Pt(y_pt))}</wp:posOffset></wp:positionV>'
        f'<wp:extent cx="{Emu(Pt(w_pt))}" cy="{Emu(Pt(h_pt))}"/><wp:effectExtent l="0" t="0" r="0" b="0"/>'
        f'<wp:wrapNone/><wp:docPr id="{z + 100}" name="{name}"/><wp:cNvGraphicFramePr/></wp:anchor>')
    anchor.append(inline.find(qn("a:graphic")))
    inline.getparent().replace(inline, anchor)


def borders(tbl, top="single", inside_h="single", inside_v="single", sz=4, color="000000"):
    tblPr = tbl._tbl.tblPr
    b = OxmlElement("w:tblBorders")
    for side, kind in (("top", top), ("left", top), ("bottom", top), ("right", top),
                       ("insideH", inside_h), ("insideV", inside_v)):
        el = OxmlElement(f"w:{side}")
        el.set(qn("w:val"), kind)
        if kind != "nil":
            el.set(qn("w:sz"), str(sz))
            el.set(qn("w:space"), "0")
            el.set(qn("w:color"), color)
        b.append(el)
    tblPr.append(b)


def table(doc, widths_pt, rows):
    t = doc.add_table(rows=rows, cols=len(widths_pt))
    t.alignment = WD_TABLE_ALIGNMENT.LEFT
    t.autofit = False
    tblPr = t._tbl.tblPr
    for tag, val in (("w:tblW", sum(widths_pt)), ("w:tblInd", 0)):
        el = OxmlElement(tag)
        el.set(qn("w:w"), twips(val))
        el.set(qn("w:type"), "dxa")
        tblPr.append(el)
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tblPr.append(layout)
    for gc, w in zip(t._tbl.tblGrid.findall(qn("w:gridCol")), widths_pt):
        gc.set(qn("w:w"), twips(w))
    for row in t.rows:
        for cell, w in zip(row.cells, widths_pt):
            cell.width = Pt(w)
    return t


def shade(cell, fill):
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    cell._tc.get_or_add_tcPr().append(shd)


def cell_text(cell, text, size=9, bold=None, italic=None, underline=None, align=None, color=None, first=True):
    p = cell.paragraphs[0] if first else cell.add_paragraph()
    return para(cell, text, size, bold, italic, underline, align, color=color, p=p)


def cell_mark(cell, size=9, bold=False, align=None):
    """Format of the (empty) paragraph mark: rows filled from {{rows:...}} copy it."""
    p = cell.paragraphs[0]
    if align is not None:
        p.alignment = align
    pPr = p._p.get_or_add_pPr()
    mark = OxmlElement("w:rPr")
    fonts = OxmlElement("w:rFonts")
    for a in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        fonts.set(qn(a), FONT)
    mark.append(fonts)
    if bold:
        mark.append(OxmlElement("w:b"))
    sz = OxmlElement("w:sz")
    sz.set(qn("w:val"), str(int(size * 2)))
    mark.append(sz)
    pPr.append(mark)
    return p


def money_table(doc, heading, note, columns, widths, field):
    p = para(doc, before=0, after=4)
    style_run(p.add_run(heading), 9.96, bold=True)
    style_run(p.add_run(" " + note), 9.96, italic=True)
    t = table(doc, widths, 3)
    borders(t)
    for cell, label in zip(t.rows[0].cells, columns):
        shade(cell, NAVY)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        cell_text(cell, label, 9, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, color="FFFFFF")
    data = t.rows[1].cells
    for i, cell in enumerate(data):
        cell_mark(cell, align=WD_ALIGN_PARAGRAPH.RIGHT if i == len(data) - 1 else None)
    style_run(data[0].paragraphs[0].add_run("{{rows:%s}}" % field), 9)
    total = t.rows[2].cells
    label = total[0].merge(total[1])
    cell_text(label, "TOTAL", 9.96, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    cell_text(total[2], "{{total:%s}}" % field, 9.96, bold=True, align=WD_ALIGN_PARAGRAPH.RIGHT)
    para(doc, size=9, after=12)


def signer(cell, role, first, space_before=14):
    para(cell, size=9, p=cell.paragraphs[0] if first else None, after=0).paragraph_format.space_before = Pt(space_before)
    cell_text(cell, "{{sig:%s}}" % role, 9, bold=True, underline=True, first=False)
    cell_text(cell, "{{title:%s}}" % role, 9, bold=True, first=False)
    cell_text(cell, "Signature over Printed Name | Date", 9, first=False)


def main():
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Pt(595.32), Pt(841.92)
    sec.left_margin, sec.right_margin = Pt(LEFT), Pt(RIGHT)
    sec.top_margin, sec.bottom_margin = Pt(64), Pt(40)
    sec.header_distance, sec.footer_distance = Pt(14), Pt(20)
    normal = doc.styles["Normal"]
    normal.font.name, normal.font.size = FONT, Pt(9)
    normal.paragraph_format.space_after = Pt(0)

    # letterhead and watermark exactly where the PDF draws them (header art repeats on every page)
    hp = sec.header.paragraphs[0]
    para(sec.header, size=2, p=hp)
    float_picture(hp, ASSETS / "watermark.png", 0, 0, 594.45, 840.75, "watermark.png", behind=True, z=1)
    float_picture(hp, ASSETS / "letterhead.png", 27.3, 14.25, 542.45, 53.05, "letterhead.png", z=2)

    body = doc.paragraphs[0] if doc.paragraphs else doc.add_paragraph()
    para(doc, "ICB-SDA-2026-007", 8, align=WD_ALIGN_PARAGRAPH.RIGHT, p=body, after=8)
    para(doc, "FINANCIAL REPORT", 14.04, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    para(doc, "as requirement for the Renewal of Recognition", 9, align=WD_ALIGN_PARAGRAPH.CENTER)
    p = para(doc, "Academic Year ", 9, align=WD_ALIGN_PARAGRAPH.CENTER, after=18)
    style_run(p.add_run("{{academic_year}}"), 9, underline=True)
    para(doc, "{{organization_name}}", 9, underline=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    para(doc, "(Name of Organization)", 9.96, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, after=14)

    t = table(doc, [248.0, CONTENT - 248.0], 3)
    borders(t)
    labels = [("Opening Balance/Fund A.Y. ", "opening_balance_ay", "opening_balance"),
              ("Closing Balance", None, "closing_balance"),
              ("Date of Submission", None, "submission_date")]
    for row, (label, ay, value) in zip(t.rows, labels):
        p = cell_text(row.cells[0], label, 9.96, bold=True)
        if ay:
            style_run(p.add_run("{{%s}}" % ay), 9.96, bold=True, underline=True)
        cell_text(row.cells[1], "{{%s}}" % value, 9.96)
    para(doc, size=9, after=18)

    money_table(doc, "SOURCES OF FUND", "(Attach a separate list for more than five (5) items)",
                ["Source", "Other Details", "Amount"], [212.5, 134.8, CONTENT - 347.3], "sources_of_fund")
    money_table(doc, "SUMMARY OF LIQUIDATION", "(Attach a separate list for more than five (5) items)",
                ["Name of Event/Expenses", "Date of Event/Payment", "Amount"], [212.5, 134.8, CONTENT - 347.3],
                "liquidation")

    # signatures: grid columns at every boundary of the PDF's two rows of boxes
    sig_w = [156.3, 77.0, 79.2, 141.8]
    t = table(doc, sig_w, 4)
    borders(t, inside_h="nil")
    for r in (0, 2):
        bar = t.rows[r].cells[0].merge(t.rows[r].cells[3])
        shade(bar, NAVY)
        cell_text(bar, "Prepared by:" if r == 0 else "NOTED BY:", 9, bold=True, color="FFFFFF")
    c = t.rows[1].cells
    auditor = c[1].merge(c[2])
    for cell, role in ((c[0], "treasurer"), (auditor, "auditor"), (c[3], "president")):
        signer(cell, role, first=True)
        para(cell, size=6)
    c = t.rows[3].cells
    left, right = c[0].merge(c[1]), c[2].merge(c[3])
    for cell, roles in ((left, ("adviser_2", "adviser_1")), (right, ("program_chair", "dean"))):
        signer(cell, roles[0], first=True)
        signer(cell, roles[1], first=False, space_before=16)
        para(cell, size=8)

    # "Received by" box for SDAO, bottom right of the page (a floating table)
    para(doc, size=9)
    t = table(doc, [129.0], 1)
    borders(t, sz=18, color=BOX_BLUE)
    pos = OxmlElement("w:tblpPr")
    for k, v in (("w:leftFromText", "180"), ("w:rightFromText", "180"), ("w:vertAnchor", "page"),
                 ("w:horzAnchor", "page"), ("w:tblpX", twips(456)), ("w:tblpY", twips(774.5))):
        pos.set(qn(k), v)
    t._tbl.tblPr.insert(0, pos)
    cell = t.rows[0].cells[0]
    p = cell_text(cell, "Received by:", 9, bold=True)
    style_run(p.add_run("______________"), 9)
    p = cell_text(cell, "Date: ", 9, bold=True, first=False)
    style_run(p.add_run("____________________"), 9)
    cell_text(cell, "(to be signed by SDAO)", 9, italic=True, align=WD_ALIGN_PARAGRAPH.CENTER, first=False)

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print("wrote", OUT)


if __name__ == "__main__":
    main()
