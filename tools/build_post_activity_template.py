"""Turn the blank SDAO Post-Activity Report (source_forms/POST-EVENT-REPORT.docx)
into the fillable template used by the generator.

The original layout, fonts and header are kept. Blank lines and signature
spaces are replaced with tokens that app/engine.py understands:

    {{field}}          inline text field
    {{sig:role}}       signatory name (from config.json or a field)
    {{line:role}}      signature underline (blank when the signer is removed)
    {{title:role}}     signatory position, editable per report
    {{annex:field}}    "attached / not attached" marker for an attachment
    {{#field}}         whole paragraph replaced by a block (paragraphs, table, images)
    {{#annexes}}       where the annex pages for attachments are appended

Run from the project root:  python tools/build_post_activity_template.py
"""
import copy
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_TAB_ALIGNMENT
from docx.shared import Inches, Twips
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
SRC = ROOT / "source_forms" / "POST-EVENT-REPORT.docx"
OUT = ROOT / "templates" / "post_activity_report" / "template.docx"


def set_text(p, text):
    """Put text in the paragraph's first run (keeping its formatting) and drop the rest."""
    runs = p.runs
    if not runs:
        p.add_run(text)
        return p.runs[0]
    runs[0].text = text
    for r in runs[1:]:
        r._r.getparent().remove(r._r)
    return runs[0]


def expect(p, starts):
    if not p.text.strip().startswith(starts):
        sys.exit(f"Unexpected template layout: wanted {starts!r}, found {p.text!r}")


def main():
    doc = Document(SRC)
    ps = doc.paragraphs

    expect(ps[0], "Top of Form")          # hidden HTML-form artifact
    expect(ps[1], "POST-ACTIVITY REPORT")
    expect(ps[3], "Name of the Organization")
    expect(ps[6], "Name of the Event")
    expect(ps[9], "Summary/Highlights")
    expect(ps[10], "Describe the what")
    expect(ps[12], "Other attachments")
    expect(ps[24], "Prepared by")
    expect(ps[30], "Reviewed by")
    expect(ps[36], "Noted by")

    # 1 & 2: blank underline -> value
    for idx, field in ((4, "organization_name"), (7, "event_name")):
        r = set_text(ps[idx], "{{%s}}" % field)
        r.bold = True

    # 3.1: instruction line -> summary block (un-numbered, indented under item 3)
    p = ps[10]
    numpr = p._p.pPr.find(qn("w:numPr"))
    if numpr is not None:
        p._p.pPr.remove(numpr)
    p.paragraph_format.left_indent = Twips(720)
    set_text(p, "{{#summary}}")

    # 4.x attachment checklist -> status marker after each item
    attachments = ["invitation", "program_flow", "committee_members", "publicity_materials",
                   "pictures", "liquidation_form", "attendance", "evaluation"]
    for i, field in enumerate(attachments):
        p = ps[13 + i]
        tpl_run = p.runs[0]
        r = copy.deepcopy(tpl_run._r)
        p._p.append(r)
        p.runs[-1].text = " {{annex:%s}}" % field
        p.runs[-1].italic = True

    # signature names on the blank line just above each underline
    for idx, left, right in ((26, "president", "secretary"),
                             (32, "adviser_1", "adviser_2"),
                             (38, "program_chair", "dean")):
        p = ps[idx]
        p.paragraph_format.tab_stops.add_tab_stop(Inches(3.0), WD_TAB_ALIGNMENT.LEFT)
        model = ps[idx + 1].runs[0]._r
        p._p.append(copy.deepcopy(model))
        r = p.runs[0]
        r.text = "{{sig:%s}}\t{{sig:%s}}" % (left, right)
        r.bold = True

    # "Prepared by": underline and position are tokens too, so each report can change the
    # officers' positions or drop the second signer (one officer preparing the report).
    for p, kind in ((ps[27], "line"), (ps[28], "title")):
        p.paragraph_format.tab_stops.add_tab_stop(Inches(3.0), WD_TAB_ALIGNMENT.LEFT)
        set_text(p, "{{%s:president}}\t{{%s:secretary}}" % (kind, kind))

    # annex pages go after the last signature line
    last = ps[40]
    new_p = copy.deepcopy(ps[39]._p)
    last._p.addnext(new_p)
    from docx.text.paragraph import Paragraph
    set_text(Paragraph(new_p, last._parent), "{{#annexes}}")

    # keep the signature section on one page so a long summary moves it whole
    for p in ps[24:40]:
        p.paragraph_format.keep_with_next = True

    # drop the hidden "Top of Form" paragraph and two of the three spacer lines
    # before "Prepared by" (gives the summary room); done last so indexes above stay valid
    for idx in (22, 23, 0):
        expect_blank = idx != 0
        if expect_blank and ps[idx].text.strip():
            sys.exit(f"Expected blank spacer at paragraph {idx}")
        ps[idx]._p.getparent().remove(ps[idx]._p)

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print("wrote", OUT)


if __name__ == "__main__":
    main()
