"""Run with:  python -m pytest -q   (from the project root)"""
import io
import json
import re

import httpx
import pytest
from docx import Document
from fastapi.testclient import TestClient

from app import main
from app.ai.base import ImageInput
from app.ai.gemini import GeminiProvider
from app.registry import get_template, load_config, merge_updates, missing_fields

NOTES = """Event: Arduino Basics Workshop
Date: October 10, 2026
Venue: Room 401
Forty-five first-year students learned breadboarding.

Program flow:
1:00 PM - Registration (Secretariat)
1:30 PM - Opening Remarks by Mr. Jhamil G. Gutierrez

Attendance:
Juan Dela Cruz, BSCpE 1
Maria Reyes, BSCpE 1
"""


@pytest.fixture
def client(tmp_path, monkeypatch):
    monkeypatch.setattr(main, "UPLOAD_DIR", tmp_path)
    monkeypatch.setenv("AI_PROVIDER", "offline")
    return TestClient(main.app)


def doc_text(data: bytes) -> str:
    d = Document(io.BytesIO(data))
    parts = [p.text for p in d.paragraphs]
    for t in d.tables:
        parts += [c.text for r in t.rows for c in r.cells]
    return "\n".join(parts)


def test_offline_chat_fills_fields(client):
    r = client.post("/api/chat", json={"template_id": "post_activity_report", "message": NOTES})
    assert r.status_code == 200, r.text
    v = r.json()["state"]["values"]
    assert v["event_name"] == "Arduino Basics Workshop"
    assert v["venue"] == "Room 401"
    assert v["program_flow"][1] == ["1:30 PM", "Opening Remarks", "Mr. Jhamil G. Gutierrez"]
    assert v["attendance"] == [["Juan Dela Cruz", "BSCpE 1"], ["Maria Reyes", "BSCpE 1"]]
    assert "breadboarding" in v["summary"]
    missing = [m["id"] for m in r.json()["missing"]]
    assert "event_name" not in missing and "evaluation" in missing


def test_docx_export_fills_tokens_and_signatories(client):
    state = {"values": {"event_name": "Arduino Basics Workshop", "summary": "Para one.\nPara two.",
                        "attendance": [["Juan Dela Cruz", "BSCpE 1"]]},
             "skipped": ["liquidation_form"], "signatories": {"president": "Juan P. Santos"}}
    r = client.post("/api/export/docx", json={"template_id": "post_activity_report", "state": state})
    assert r.status_code == 200
    text = doc_text(r.content)
    assert "{{" not in text
    assert "Academic Society of Computer Engineering Discipline (ASCEnD)" in text  # default
    assert "Juan P. Santos" in text and "Dr. Corazon B. Rebong, PCpE" in text
    assert "Engr. Jhamil G. Gutierrez, MSCpE" in text and "Engr. Mary Jane F. Somao-i, MSCpE" in text
    assert "List of Attendance – see Annex A" in text
    assert "Liquidation Form – N/A" in text
    assert "Para two." in text and "Juan Dela Cruz" in text


def test_preview_and_skip(client):
    state = {"values": {}, "skipped": ["invitation"]}
    r = client.post("/api/preview", json={"template_id": "post_activity_report", "state": state})
    body = r.json()
    assert 'data-field="event_name"' in body["html"]
    assert 'data-sig="dean"' in body["html"]
    assert "invitation" not in [m["id"] for m in body["missing"]]


def test_image_upload_goes_into_annex(client):
    import base64
    png = base64.b64decode("iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mNk+M9QDwADhgGAWjR9awAAAABJRU5ErkJggg==")
    up = client.post("/api/uploads", files=[("files", ("p.png", png, "image/png"))]).json()
    state = {"values": {"pictures": [{"id": up[0]["id"], "caption": "Group photo"}]}}
    r = client.post("/api/export/docx", json={"template_id": "post_activity_report", "state": state})
    d = Document(io.BytesIO(r.content))
    assert len(d.inline_shapes) == 1
    assert "Figure 1. Group photo" in doc_text(r.content)


def _docx_bytes(text: str) -> bytes:
    d = Document()
    for line in text.splitlines():
        d.add_paragraph(line)
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def test_document_upload_feeds_chat(client):
    docx_mime = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    up = client.post("/api/uploads", files=[("files", ("notes.docx", _docx_bytes(NOTES), docx_mime))])
    assert up.status_code == 200, up.text
    rec = up.json()[0]
    assert rec["kind"] == "document" and rec["chars"] > 0
    r = client.post("/api/chat", json={"template_id": "post_activity_report", "file_ids": [rec["id"]]})
    v = r.json()["state"]["values"]
    assert v["event_name"] == "Arduino Basics Workshop" and v["venue"] == "Room 401"
    assert ["Juan Dela Cruz", "BSCpE 1"] in v["attendance"]


def test_pdf_upload_renders_pages_but_not_into_pictures(client):
    pdf = open("examples/Post-Activity Report - Arduino Basics Workshop 2026.pdf", "rb").read()
    rec = client.post("/api/uploads", files=[("files", ("report.pdf", pdf, "application/pdf"))]).json()[0]
    assert rec["pages"] and rec["chars"] > 0
    assert client.get(rec["pages"][0]["url"]).headers["content-type"] == "image/png"
    r = client.post("/api/chat", json={"template_id": "post_activity_report", "file_ids": [rec["id"]]})
    assert "pictures" not in r.json()["state"]["values"]


def test_unsupported_upload_is_rejected(client):
    r = client.post("/api/uploads", files=[("files", ("tool.exe", b"MZ", "application/octet-stream"))])
    assert r.status_code == 415


def test_merge_updates_validates_types():
    tpl, cfg = get_template("post_activity_report"), load_config()
    state, changed = merge_updates(tpl, {"skipped": ["program_flow"]}, {
        "program_flow": [{"Time": "8:00", "Activity": "Opening"}], "not_a_field": "x", "event_name": 42})
    assert state["values"]["program_flow"] == [["8:00", "Opening", ""]]
    assert state["values"]["event_name"] == "42"
    assert "program_flow" not in state["skipped"]
    assert "not_a_field" not in state["values"]


def test_gemini_provider_parses_json(monkeypatch):
    seen = {}

    def fake_post(url, json=None, timeout=None, headers=None):
        seen["body"] = json
        reply = {"updates": {"event_name": "Hackathon"}, "reply": "Got it."}
        return httpx.Response(200, json={"candidates": [{"content": {"parts": [{"text": __import__("json").dumps(reply)}]}}]})

    monkeypatch.setattr(httpx, "post", fake_post)
    tpl, cfg = get_template("post_activity_report"), load_config()
    p = GeminiProvider(api_key="test")
    res = p.extract(tpl=tpl, values={}, skipped=[], missing=missing_fields(tpl, {}, cfg), message="notes",
                    images=[ImageInput(id="abc", mime="image/png", data=b"x")], history=[])
    assert res.updates == {"event_name": "Hackathon"} and res.reply == "Got it."
    parts = seen["body"]["contents"][-1]["parts"]
    assert any("inline_data" in part for part in parts)
    assert "program_flow" in seen["body"]["systemInstruction"]["parts"][0]["text"]


def test_password_protects_everything_but_health(client, monkeypatch):
    monkeypatch.setattr(main, "PASSWORD", "s3cret")
    assert client.get("/api/templates").status_code == 401
    assert client.get("/api/templates", auth=("any", "wrong")).status_code == 401
    assert client.get("/api/templates", auth=("any", "s3cret")).status_code == 200
    assert client.get("/api/health").status_code == 200


def test_users_own_key_is_used_for_chat(client, monkeypatch):
    seen = {}

    def fake_post(url, json=None, timeout=None, headers=None):
        seen["key"] = headers["x-goog-api-key"]
        if headers["x-goog-api-key"] == "bad":
            return httpx.Response(400, text='{"error": {"status": "INVALID_ARGUMENT", "details": [{"reason": "API_KEY_INVALID"}]}}')
        reply = {"updates": {"event_name": "Hackathon"}, "reply": "Got it."}
        return httpx.Response(200, json={"candidates": [{"content": {"parts": [{"text": __import__("json").dumps(reply)}]}}]})

    monkeypatch.setattr(httpx, "post", fake_post)
    body = {"template_id": "post_activity_report", "message": "notes"}
    r = client.post("/api/chat", json=body, headers={"X-Gemini-Key": "user-key"})
    assert r.json()["provider"] == "gemini" and seen["key"] == "user-key"
    r = client.post("/api/chat", json=body, headers={"X-Gemini-Key": "bad"})
    assert r.status_code == 502 and "rejected the API key" in r.json()["detail"]
    assert client.post("/api/chat", json=body).json()["provider"] == "offline"  # no key: server default


def _prepared_by_lines(data: bytes) -> list[str]:
    ps = [p.text for p in Document(io.BytesIO(data)).paragraphs]
    i = ps.index("Prepared by")
    return ps[i + 2:i + 5]  # names, underlines, positions


def test_prepared_by_defaults_to_president_and_secretary(client):
    state = {"signatories": {"president": "Juan P. Santos", "secretary": "Ana L. Cruz"}}
    r = client.post("/api/export/docx", json={"template_id": "post_activity_report", "state": state})
    names, lines, titles = _prepared_by_lines(r.content)
    assert names == "Juan P. Santos\tAna L. Cruz"
    assert lines.count("_") == 60 and titles == "President\tSecretary"


def test_prepared_by_one_officer_with_custom_position(client):
    state = {"signatories": {"president": "Maria Reyes"}, "sig_titles": {"president": "Vice President - Internal"},
             "sig_removed": ["secretary"]}
    body = {"template_id": "post_activity_report", "state": state}
    names, lines, titles = _prepared_by_lines(client.post("/api/export/docx", json=body).content)
    assert names.strip() == "Maria Reyes" and "SECRETARY" not in names
    assert lines.strip() == "_" * 30 and titles.strip() == "Vice President - Internal"
    html = client.post("/api/preview", json=body).json()["html"]
    assert 'data-sig-add="secretary"' in html and "Vice President - Internal" in html


def test_only_removable_signatories_can_be_removed(client):
    state = {"sig_removed": ["dean"]}
    r = client.post("/api/export/docx", json={"template_id": "post_activity_report", "state": state})
    assert "Dr. Corazon B. Rebong, PCpE" in doc_text(r.content)


PROPOSAL = {"values": {
    "event_name": "Arduino Basics Workshop", "event_date": "October 24, 2026",
    "introduction": "Para one.\nPara two.",
    "activity_overview": [["Talk", "Learn", "Lecture", "Projector", "Juan"], ["Lab", "Build", "Hands-on", "Kits", "Ana"]],
    "budget_items": [["Snacks", "45", "₱50.00", "₱2,250.00"], ["Kits", "10", "₱300.00", "₱3,000.00"]],
    "program_flow": [["Oct 24, 1:00 PM", "Room 401", "Registration", "Secretariat"]]},
    "signatories": {"president": "Maria Reyes"}}


def test_proposal_fills_form_tables_in_place(client):
    r = client.post("/api/export/docx", json={"template_id": "activity_proposal", "state": PROPOSAL})
    assert r.status_code == 200
    d = Document(io.BytesIO(r.content))
    text = doc_text(r.content)
    assert "{{" not in text
    overview, budget, flow, approval = d.tables[1], d.tables[2], d.tables[3], d.tables[4]
    assert [c.text for c in overview.rows[2].cells] == ["Lab", "Build", "Hands-on", "Kits", "Ana"]
    assert len(overview.rows) == 3 and len(budget.rows) == 4  # header + rows (+ TOTAL)
    assert budget.rows[-1].cells[-1].text == "₱5,250.00"
    assert flow.rows[0].cells[0].text == "TITLE: Arduino Basics Workshop"  # defaults to the event title
    signers = approval.rows[1].cells
    assert "Maria Reyes\nPresident" in signers[0].text
    assert "Engr. Jhamil G. Gutierrez, MSCpE\nRSO Adviser" in signers[1].text
    # the yellow header styling of the form's table is kept
    assert 'w:fill="FFC000"' in overview.rows[0].cells[0]._tc.xml


def test_proposal_empty_table_keeps_one_blank_row(client):
    r = client.post("/api/export/docx", json={"template_id": "activity_proposal", "state": {}})
    d = Document(io.BytesIO(r.content))
    assert len(d.tables[1].rows) == 2 and not any(c.text for c in d.tables[1].rows[1].cells)
    assert d.tables[2].rows[-1].cells[-1].text == ""  # no amounts, no total


def test_proposal_preview_has_form_art_and_editable_rows(client):
    html = client.post("/api/preview", json={"template_id": "activity_proposal", "state": PROPOSAL}).json()["html"]
    assert html.count('class="dfloat') == 5 and 'class="dpic crop"' in html  # header/footer art, cropped logo
    assert 'data-table="budget_items"' in html and 'data-field="budget_items" data-row="1" data-col="3"' in html
    assert "₱5,250.00" in html and "{{" not in html
    url = re.search(r'/api/templates/activity_proposal/media/[\w.-]+', html).group(0)
    assert client.get(url).headers["content-type"] == "image/png"
    assert client.get("/api/templates/activity_proposal/media/..%2Fdocument.xml").status_code in (400, 404)


def test_prompt_uses_each_forms_own_rules():
    from app.ai.prompt import system_prompt
    proposal = system_prompt(get_template("activity_proposal"), {}, [], [], [])
    report = system_prompt(get_template("post_activity_report"), {}, [], [], [])
    assert "future tense" in proposal and '"summary"' not in proposal
    assert 'Write "summary"' in report and "future tense" not in report
